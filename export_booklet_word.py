# -*- coding: utf-8 -*-
"""
山陰山陽 13 天旅行小冊子 Word (.docx) 產生器
每日前後自成獨立篇章，具備強制分頁、景點詳細資訊、精準車次時刻、車資預算表及紀念章蓋印框。
"""

import os
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """設定儲存格背景底色"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """設定儲存格內部 padding (單位: dxa, 20 dxa = 1 pt)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    設定儲存格框線
    kwargs: top, bottom, left, right 各自為 dict(sz=12, val='single', color='CBD5E1')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name, border_props in kwargs.items():
        node = OxmlElement(f'w:{border_name}')
        for key, val in border_props.items():
            node.set(qn(f'w:{key}'), str(val))
        tcBorders.append(node)
    tcPr.append(tcBorders)

def add_header_footer(doc):
    """設定全份文件頁首與頁尾"""
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.different_first_page_header_footer = True
        
        # 頁首
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("山陰山陽壯遊 13 天・旅行隨身手冊")
        hrun.font.name = "Microsoft JhengHei"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(148, 163, 184)
        
        # 頁尾
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("— 旅行隨身手冊 —")
        frun.font.name = "Microsoft JhengHei"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(148, 163, 184)

def format_run(run, font_name="Microsoft JhengHei", size_pt=10.5, bold=False, color_rgb=(51, 65, 85), italic=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def create_booklet_word():
    # 讀取 JSON 資料
    with open('spots_data.json', 'r', encoding='utf-8') as f:
        data = json.load(f)

    spots_map = {s['id']: s for s in data.get('spots', [])}
    itinerary = data.get('itinerary', [])
    fares_list = data.get('fares', [])

    doc = Document()
    add_header_footer(doc)

    # ==========================================
    # 1. 封面 (Cover Page)
    # ==========================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(40)
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub0 = p_pre.add_run("🇯🇵 2027 嚴選壯遊隨身專屬手冊")
    format_run(r_sub0, size_pt=13, bold=True, color_rgb=(37, 99, 235))

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("山陰山陽 13 天縱走\n名城・名湯・絕景漫遊手冊")
    format_run(r_title, size_pt=24, bold=True, color_rgb=(30, 58, 138))

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(36)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("日本100名城巡禮 ｜ 三大名湯 ｜ 日本海極品松葉蟹 ｜ 世界遺產石見銀山")
    format_run(r_sub, size_pt=11, color_rgb=(100, 116, 139))

    # 封面資訊卡片 (表格)
    card_table = doc.add_table(rows=6, cols=2)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_table.autofit = False

    info_data = [
        ("📅 旅行日期", "2027 年 1 月 13 日 (三) ～ 1 月 25 日 (一) 共 13 天 12 夜"),
        ("👥 隨行成員", "獅子 & 方 (雙人自助旅行・全程大眾運輸)"),
        ("✈️ 國際航班", "台灣虎航直飛岡山\n去程 IT214：01/13 11:30 TPE ➔ 15:05 OKJ\n回程 IT215：01/25 15:55 OKJ ➔ 17:40 TPE"),
        ("🎫 交通周遊券", "JR 山陰&山陽鐵路周遊券 7日券\n【Day 3 ～ Day 9 正式啟用】含山陰特急/丹後鐵道/智頭急行指定席"),
        ("🏯 攻城重點", "鬼之城、津山城、鳥取城、月山富田城、松江城(國寶)、津和野城、萩城、岩國城、廣島城、吉田郡山城、福山城、備中松山城"),
        ("🆘 緊急聯絡", "日本報警：110 ｜ 救護火警：119\n外交部旅外急難求助：+886-800-085-095\n台北駐大阪經濟文化辦事處：+81-6-6227-8623")
    ]

    col_widths = [Inches(1.8), Inches(4.7)]
    for row_idx, (k, v) in enumerate(info_data):
        row = card_table.rows[row_idx]
        
        # 標題格
        c0 = row.cells[0]
        c0.width = col_widths[0]
        set_cell_background(c0, "F8FAFC")
        set_cell_margins(c0, top=140, bottom=140, left=180, right=140)
        set_cell_border(c0, bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                            top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                            left={'sz': 12, 'val': 'single', 'color': '2563EB'},
                            right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(k)
        format_run(r0, size_pt=9.5, bold=True, color_rgb=(30, 58, 138))

        # 內容格
        c1 = row.cells[1]
        c1.width = col_widths[1]
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c1, top=140, bottom=140, left=180, right=180)
        set_cell_border(c1, bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                            top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                            left={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                            right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(v)
        format_run(r1, size_pt=9.5, color_rgb=(51, 65, 85))

    doc.add_page_break()

    # ==========================================
    # 2. 13天行程總覽速查表 (Overview)
    # ==========================================
    p_ov_title = doc.add_paragraph()
    p_ov_title.paragraph_format.space_before = Pt(10)
    p_ov_title.paragraph_format.space_after = Pt(12)
    r_ov_t = p_ov_title.add_run("📋 13 天行程大綱與夜宿速查表")
    format_run(r_ov_t, size_pt=16, bold=True, color_rgb=(30, 58, 138))

    ov_table = doc.add_table(rows=1, cols=4)
    ov_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ov_table.autofit = False
    
    headers = ["天數", "日期 / 樞紐站", "當日主題核心精華", "夜宿地點 / 旅館類型"]
    widths = [Inches(0.8), Inches(1.5), Inches(2.6), Inches(1.6)]

    hdr_cells = ov_table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_border(hdr_cells[i], bottom={'sz': 8, 'val': 'single', 'color': '1E3A8A'})
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        format_run(r, size_pt=9.5, bold=True, color_rgb=(255, 255, 255))

    for day in itinerary:
        d_num = day.get('dayNumber', '')
        d_date = day.get('date', '')
        d_base = day.get('baseStation', '').split(' / ')[0]
        d_title = day.get('title', '').replace(f"Day {d_num}：", "")
        d_lodge = day.get('lodging', '')

        row = ov_table.add_row()
        cells = row.cells
        
        bg_col = "F8FAFC" if d_num % 2 == 1 else "FFFFFF"
        for i in range(4):
            cells[i].width = widths[i]
            set_cell_background(cells[i], bg_col)
            set_cell_margins(cells[i], top=100, bottom=100, left=120, right=120)
            set_cell_border(cells[i], bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                      top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                      left={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                      right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
        
        # 天數
        p0 = cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(f"D{d_num}")
        format_run(r0, size_pt=9.5, bold=True, color_rgb=(37, 99, 235))

        # 日期
        p1 = cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(f"{d_date}\n({d_base})")
        format_run(r1, size_pt=8.5, color_rgb=(71, 85, 105))

        # 標題
        p2 = cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p2.add_run(d_title)
        format_run(r2, size_pt=9, bold=True, color_rgb=(30, 58, 138))

        # 住宿
        p3 = cells[3].paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r3 = p3.add_run(d_lodge)
        format_run(r3, size_pt=8.5, color_rgb=(51, 65, 85))

    doc.add_page_break()

    # ==========================================
    # 3. 每日獨立篇章 (每天自成一個段落，強制分頁)
    # ==========================================
    for day_idx, day in enumerate(itinerary):
        d_num = day.get('dayNumber', '')
        d_date = day.get('date', '')
        d_title = day.get('title', '')
        d_base = day.get('baseStation', '')
        d_lodging = day.get('lodging', '')
        d_loc = day.get('lodgingLocation', '')
        d_recom = day.get('recommendedHotels', '')
        d_reason = day.get('lodgingReason', '')
        d_notes = day.get('notes', '')
        spot_ids = day.get('spotIds', [])

        # 每一天的醒目主標題
        p_day_title = doc.add_paragraph()
        p_day_title.paragraph_format.space_before = Pt(8)
        p_day_title.paragraph_format.space_after = Pt(6)
        r_day_t = p_day_title.add_run(f"{d_title}")
        format_run(r_day_t, size_pt=14, bold=True, color_rgb=(30, 58, 138))

        # 當日基本資訊卡片
        base_table = doc.add_table(rows=3, cols=2)
        base_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        base_table.autofit = False
        b_widths = [Inches(1.5), Inches(5.0)]

        pass_status = "🎫 JR Pass 7日券【適用日】(全額免費搭乘指定席)" if 3 <= d_num <= 9 else "💳 刷交通 IC 卡 / 現金 (非 Pass 天數)"

        base_rows_data = [
            ("🗓️ 日期與樞紐", f"{d_date} ｜ 樞紐站：{d_base} ｜ 交通狀態：{pass_status}"),
            ("🏨 夜宿地點", f"{d_lodging} ({d_loc})\n推薦飯店：{d_recom}"),
            ("💡 入住理由與安排", f"{d_reason}")
        ]

        for r_i, (k, v) in enumerate(base_rows_data):
            row = base_table.rows[r_i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width, c1.width = b_widths[0], b_widths[1]

            set_cell_background(c0, "F1F5F9")
            set_cell_margins(c0, top=80, bottom=80, left=120, right=100)
            set_cell_border(c0, bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                left={'sz': 10, 'val': 'single', 'color': '2563EB'},
                                right={'sz': 4, 'val': 'single', 'color': 'CBD5E1'})
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(k)
            format_run(r0, size_pt=9, bold=True, color_rgb=(30, 58, 138))

            set_cell_background(c1, "FFFFFF")
            set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
            set_cell_border(c1, bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                left={'sz': 4, 'val': 'single', 'color': 'CBD5E1'},
                                right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(v)
            format_run(r1, size_pt=9, color_rgb=(51, 65, 85))

        # ------------------------------------------
        # 景點明細表
        # ------------------------------------------
        p_sp_hdr = doc.add_paragraph()
        p_sp_hdr.paragraph_format.space_before = Pt(12)
        p_sp_hdr.paragraph_format.space_after = Pt(4)
        r_sph = p_sp_hdr.add_run(f"📍 Day {d_num} 景點動線與蓋章資訊")
        format_run(r_sph, size_pt=11, bold=True, color_rgb=(37, 99, 235))

        spots_for_day = [spots_map.get(sid) for sid in spot_ids if sid in spots_map]

        if spots_for_day:
            sp_table = doc.add_table(rows=1, cols=5)
            sp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            sp_table.autofit = False
            sp_widths = [Inches(0.5), Inches(1.8), Inches(1.5), Inches(1.3), Inches(1.4)]
            sp_headers = ["序", "景點名稱", "蓋章處 / 類別", "開放時間 / 公休", "交通接駁與提示"]

            for i, h in enumerate(sp_headers):
                c = sp_table.rows[0].cells[i]
                c.width = sp_widths[i]
                set_cell_background(c, "2563EB")
                set_cell_margins(c, top=80, bottom=80, left=100, right=100)
                p = c.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(h)
                format_run(r, size_pt=8.5, bold=True, color_rgb=(255, 255, 255))

            for s_idx, sp in enumerate(spots_for_day, 1):
                s_name = sp.get('name', '')
                s_pref = sp.get('prefecture', '')
                s_stamp = sp.get('stampLocation', '無')
                s_open = sp.get('openHours', '全年無休')
                s_closed = sp.get('closedDayText', '')
                s_transit = sp.get('transitInfo', '')

                row = sp_table.add_row()
                row_cells = row.cells
                bg = "F8FAFC" if s_idx % 2 == 1 else "FFFFFF"

                for ci in range(5):
                    row_cells[ci].width = sp_widths[ci]
                    set_cell_background(row_cells[ci], bg)
                    set_cell_margins(row_cells[ci], top=70, bottom=70, left=90, right=90)
                    set_cell_border(row_cells[ci], bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                   top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                   left={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                   right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})

                # 序號
                p_c0 = row_cells[0].paragraphs[0]
                p_c0.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c0 = p_c0.add_run(str(s_idx))
                format_run(r_c0, size_pt=9, bold=True, color_rgb=(30, 58, 138))

                # 景點名稱
                p_c1 = row_cells[1].paragraphs[0]
                r_c1 = p_c1.add_run(f"{s_name}\n({s_pref})")
                format_run(r_c1, size_pt=9, bold=True, color_rgb=(30, 58, 138))

                # 蓋章處
                p_c2 = row_cells[2].paragraphs[0]
                r_c2 = p_c2.add_run(f"印章：{s_stamp}" if s_stamp and s_stamp != '無' else "一般參觀")
                format_run(r_c2, size_pt=8, color_rgb=(217, 119, 6) if s_stamp and s_stamp != '無' else (71, 85, 105))

                # 開放時間
                p_c3 = row_cells[3].paragraphs[0]
                r_c3 = p_c3.add_run(f"{s_open}\n休：{s_closed if s_closed else '無休'}")
                format_run(r_c3, size_pt=8, color_rgb=(51, 65, 85))

                # 交通
                p_c4 = row_cells[4].paragraphs[0]
                r_c4 = p_c4.add_run(s_transit[:75] + ("..." if len(s_transit) > 75 else ""))
                format_run(r_c4, size_pt=8, color_rgb=(71, 85, 105))

        # ------------------------------------------
        # 當日交通叮嚀與時刻備忘
        # ------------------------------------------
        p_nt_hdr = doc.add_paragraph()
        p_nt_hdr.paragraph_format.space_before = Pt(10)
        p_nt_hdr.paragraph_format.space_after = Pt(4)
        r_nth = p_nt_hdr.add_run("⏱️ 詳細交通時刻銜接與行程叮嚀")
        format_run(r_nth, size_pt=11, bold=True, color_rgb=(37, 99, 235))

        if d_notes:
            notes_lines = [l.strip() for l in d_notes.split('\n') if l.strip()]
            for line in notes_lines:
                p_note = doc.add_paragraph()
                p_note.paragraph_format.space_before = Pt(2)
                p_note.paragraph_format.space_after = Pt(2)
                p_note.paragraph_format.left_indent = Inches(0.15)
                
                # 特殊高亮
                r_n = p_note.add_run(line)
                if "JR Pass" in line:
                    format_run(r_n, size_pt=8.5, bold=True, color_rgb=(37, 99, 235))
                elif "免計程車" in line or "07:50" in line or "特急" in line:
                    format_run(r_n, size_pt=8.5, bold=True, color_rgb=(15, 23, 42))
                else:
                    format_run(r_n, size_pt=8.5, color_rgb=(51, 65, 85))

        # ------------------------------------------
        # 當日車資與自費預算表
        # ------------------------------------------
        day_fares = [f for f in fares_list if f.get('dayNumber') == d_num]
        if day_fares:
            p_fare_hdr = doc.add_paragraph()
            p_fare_hdr.paragraph_format.space_before = Pt(10)
            p_fare_hdr.paragraph_format.space_after = Pt(4)
            r_fh = p_fare_hdr.add_run("💴 當日車資明細表 (Pass 覆蓋 vs 自費)")
            format_run(r_fh, size_pt=11, bold=True, color_rgb=(37, 99, 235))

            f_table = doc.add_table(rows=1, cols=5)
            f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            f_table.autofit = False
            f_widths = [Inches(2.5), Inches(0.9), Inches(0.8), Inches(1.1), Inches(1.2)]
            f_headers = ["乘車區間 / 路線", "車種", "金額(円)", "Pass 狀態", "乘車備註"]

            for fi, fh in enumerate(f_headers):
                fc = f_table.rows[0].cells[fi]
                fc.width = f_widths[fi]
                set_cell_background(fc, "334155")
                set_cell_margins(fc, top=60, bottom=60, left=80, right=80)
                fp = fc.paragraphs[0]
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fr = fp.add_run(fh)
                format_run(fr, size_pt=8, bold=True, color_rgb=(255, 255, 255))

            total_self_pay = 0
            for fare in day_fares:
                f_title = fare.get('title', '')
                f_type = fare.get('transitTypeName', '')
                f_amt = fare.get('amount', 0)
                f_pass = fare.get('passCovered', False)
                f_pass_note = fare.get('passNote', '')
                f_note = fare.get('notes', '')

                if not f_pass:
                    total_self_pay += f_amt

                frow = f_table.add_row()
                fcells = frow.cells
                for fci in range(5):
                    fcells[fci].width = f_widths[fci]
                    set_cell_background(fcells[fci], "FFFFFF")
                    set_cell_margins(fcells[fci], top=50, bottom=50, left=70, right=70)
                    set_cell_border(fcells[fci], bottom={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                 top={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                 left={'sz': 4, 'val': 'single', 'color': 'E2E8F0'},
                                                 right={'sz': 4, 'val': 'single', 'color': 'E2E8F0'})

                # 區間
                p_f0 = fcells[0].paragraphs[0]
                r_f0 = p_f0.add_run(f_title)
                format_run(r_f0, size_pt=8, bold=True, color_rgb=(30, 58, 138))

                # 車種
                p_f1 = fcells[1].paragraphs[0]
                p_f1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_f1 = p_f1.add_run(f_type)
                format_run(r_f1, size_pt=8, color_rgb=(71, 85, 105))

                # 金額
                p_f2 = fcells[2].paragraphs[0]
                p_f2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r_f2 = p_f2.add_run(f"¥{f_amt:,}")
                format_run(r_f2, size_pt=8, bold=True, color_rgb=(220, 38, 38) if not f_pass else (16, 185, 129))

                # Pass 狀態
                p_f3 = fcells[3].paragraphs[0]
                p_f3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_f3 = p_f3.add_run("Pass 免費" if f_pass else "⚠️ 需自費")
                format_run(r_f3, size_pt=8, bold=True, color_rgb=(16, 185, 129) if f_pass else (217, 119, 6))

                # 備註
                p_f4 = fcells[4].paragraphs[0]
                r_f4 = p_f4.add_run(f_note[:60])
                format_run(r_f4, size_pt=7.5, color_rgb=(100, 116, 139))

        # ------------------------------------------
        # 小冊子專用：紀念章戳蓋印處與手寫記事區
        # ------------------------------------------
        p_stamp_hdr = doc.add_paragraph()
        p_stamp_hdr.paragraph_format.space_before = Pt(12)
        p_stamp_hdr.paragraph_format.space_after = Pt(4)
        r_sth = p_stamp_hdr.add_run("💮 隨身手冊專區：日本100名城蓋章印框 ＆ 隨行筆記")
        format_run(r_sth, size_pt=10, bold=True, color_rgb=(100, 116, 139))

        stamp_table = doc.add_table(rows=1, cols=2)
        stamp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        stamp_table.autofit = False
        st_widths = [Inches(2.4), Inches(4.1)]

        # 左側：蓋章方框
        st_c0 = stamp_table.rows[0].cells[0]
        st_c0.width = st_widths[0]
        set_cell_background(st_c0, "FAFAFA")
        set_cell_margins(st_c0, top=140, bottom=140, left=140, right=140)
        set_cell_border(st_c0, bottom={'sz': 12, 'val': 'dashed', 'color': '94A3B8'},
                               top={'sz': 12, 'val': 'dashed', 'color': '94A3B8'},
                               left={'sz': 12, 'val': 'dashed', 'color': '94A3B8'},
                               right={'sz': 12, 'val': 'dashed', 'color': '94A3B8'})
        p_st0 = st_c0.paragraphs[0]
        p_st0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_st0 = p_st0.add_run("【 百名城 / 紀念印章蓋印處 】\n\n\n\n(預留 6cm x 6cm 印章空間)")
        format_run(r_st0, size_pt=8.5, color_rgb=(148, 163, 184))

        # 右側：手寫筆記橫線
        st_c1 = stamp_table.rows[0].cells[1]
        st_c1.width = st_widths[1]
        set_cell_background(st_c1, "FFFFFF")
        set_cell_margins(st_c1, top=100, bottom=100, left=140, right=140)
        set_cell_border(st_c1, bottom={'sz': 6, 'val': 'single', 'color': 'CBD5E1'},
                               top={'sz': 6, 'val': 'single', 'color': 'CBD5E1'},
                               left={'sz': 6, 'val': 'single', 'color': 'CBD5E1'},
                               right={'sz': 6, 'val': 'single', 'color': 'CBD5E1'})
        p_st1 = st_c1.paragraphs[0]
        r_st1 = p_st1.add_run("📝 當日心得 / 推薦美食記帳 / 票根貼附：\n\n\n\n\n")
        format_run(r_st1, size_pt=8.5, color_rgb=(148, 163, 184))

        # 每一天的結尾：如果不是最後一天，強制分頁！
        if day_idx < len(itinerary) - 1:
            doc.add_page_break()

    # 儲存檔案
    output_filename = "山陰山陽13天旅行小冊子.docx"
    doc.save(output_filename)
    print(f"成功產生 Word 小冊子：{output_filename}")

if __name__ == "__main__":
    create_booklet_word()
